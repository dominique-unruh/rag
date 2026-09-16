"""The ``pptx`` translator (PLAN §5.4).

Converts a PowerPoint OOXML presentation (``.pptx``/``.potx``/``.ppsx``) to
Markdown using the ``python-pptx`` library: one ``## Slide N`` section per slide,
gathering each slide's shape text (recursing into grouped shapes), any tables,
and the speaker notes. Slides that carry no text at all are skipped, so numbering
gaps honestly reflect empty slides.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..base import Document, Translation, Translator, make_translation

# The content type declared for the *main* presentation part inside the OOXML
# package's ``[Content_Types].xml``. python-pptx only accepts the deck spelling,
# so a template (.potx) or slideshow (.ppsx) — structurally identical decks that
# merely carry a different main content type — is rejected outright. We rewrite
# the template/slideshow spelling to the deck spelling before opening.
_DECK_MAIN_CT: str = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
)
_TEMPLATE_MAIN_CT: str = (
    "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
)
_SLIDESHOW_MAIN_CT: str = (
    "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml"
)
# The package part that maps part names to content types.
_CONTENT_TYPES_PART: str = "[Content_Types].xml"


def _as_deck(content: bytes) -> bytes:
    """Return ``content`` with any template/slideshow main content type made a deck.

    python-pptx refuses to open a .potx/.ppsx because its main part is typed as a
    template/slideshow rather than a presentation. The parts themselves are the
    same, so we patch just the content-type declaration and re-pack the zip. A
    genuine deck contains neither template nor slideshow spelling and is returned
    unchanged (no re-zip).
    """
    source = zipfile.ZipFile(io.BytesIO(content))
    content_types = source.read(_CONTENT_TYPES_PART).decode("utf-8")
    # Nothing to do for an ordinary deck (or macro deck): leave the bytes as-is.
    if _TEMPLATE_MAIN_CT not in content_types and _SLIDESHOW_MAIN_CT not in content_types:
        return content
    patched = content_types.replace(_TEMPLATE_MAIN_CT, _DECK_MAIN_CT).replace(
        _SLIDESHOW_MAIN_CT, _DECK_MAIN_CT
    )
    out = io.BytesIO()
    # Rebuild the archive, swapping only the content-types part; every other
    # member is copied verbatim (preserving its per-entry compression).
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dest:
        for item in source.infolist():
            data = (
                patched.encode("utf-8")
                if item.filename == _CONTENT_TYPES_PART
                else source.read(item.filename)
            )
            dest.writestr(item, data)
    return out.getvalue()


def _cell(value: object) -> str:
    """Render one table cell as Markdown-table-safe text.

    Pipes are escaped (they are the table delimiter) and any newlines are
    flattened so one table row stays on a single Markdown line.
    """
    text = str(value)
    text = text.replace("|", "\\|")
    return text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")


def _iter_leaf_shapes(shapes: object) -> Iterator[object]:
    """Yield every leaf shape under ``shapes``, descending into group shapes.

    A slide's shape tree can nest groups arbitrarily; python-pptx exposes a
    group's children under ``shape.shapes``. Flattening here lets the text/table
    extraction below treat groups transparently.
    """
    for shape in shapes:  # type: ignore[attr-defined]
        # Recurse into groups; their own bounding shape carries no text of interest.
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_leaf_shapes(shape.shapes)
        else:
            yield shape


def _table_to_markdown(table: object) -> str:
    """Render a pptx table as a Markdown table.

    The first row is used as the header (Markdown tables require a header +
    separator even when the slide table has no real header row).
    """
    # Materialise every row as a list of cleaned cell strings.
    rows: list[list[str]] = [
        [_cell(cell.text) for cell in row.cells] for row in table.rows  # type: ignore[attr-defined]
    ]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    # Pad short rows so every Markdown row has the same column count.
    norm = [(row + [""] * width)[:width] for row in rows]
    lines = [
        "| " + " | ".join(norm[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines += ["| " + " | ".join(row) + " |" for row in norm[1:]]
    return "\n".join(lines)


def _text_frame_lines(text_frame: object) -> list[str]:
    """Return the non-empty paragraph lines of a text frame.

    Each paragraph becomes one line; ``paragraph.text`` already concatenates its
    runs. Blank paragraphs (spacing) are dropped so they don't pad the output.
    """
    lines: list[str] = []
    for paragraph in text_frame.paragraphs:  # type: ignore[attr-defined]
        line = paragraph.text.strip()
        if line:
            lines.append(line)
    return lines


def _shape_block(shape: object) -> str | None:
    """Render one leaf shape to a Markdown block, or ``None`` if it has no text.

    Tables are rendered as Markdown tables; any other shape with a text frame
    contributes its paragraph lines. Picture/media/empty shapes yield ``None``.
    """
    # Tables first: a table shape also reports has_text_frame in some builds, but
    # its cell text lives in the table, not a frame.
    if shape.has_table:
        return _table_to_markdown(shape.table) or None
    if shape.has_text_frame:
        lines = _text_frame_lines(shape.text_frame)
        return "\n".join(lines) if lines else None
    return None


def _notes_text(slide: object) -> str | None:
    """Return the slide's speaker-notes text, or ``None`` if there are none.

    ``has_notes_slide`` avoids creating an empty notes slide as a side effect of
    the check (accessing ``notes_slide`` would materialise one).
    """
    if not slide.has_notes_slide:
        return None
    text = slide.notes_slide.notes_text_frame.text.strip()
    return text or None


def _slide_markdown(index: int, slide: object) -> str | None:
    """Render one slide (1-based ``index``) to a Markdown section, or ``None``.

    Collects every leaf shape's text block in reading order, then appends a
    ``### Notes`` block for any speaker notes. Returns ``None`` when the slide has
    no text at all, so empty slides are skipped rather than emitting a bare header.
    """
    blocks: list[str] = []
    for shape in _iter_leaf_shapes(slide.shapes):
        block = _shape_block(shape)
        if block:
            blocks.append(block)
    notes = _notes_text(slide)
    if notes:
        blocks.append(f"### Notes\n\n{notes}")
    if not blocks:
        return None
    # Blank line between the heading and each block keeps Markdown structure clean.
    return f"## Slide {index}\n\n" + "\n\n".join(blocks)


class Pptx(Translator):
    """Converts a PowerPoint OOXML presentation to Markdown, one section per slide."""

    def translate(self, document: Document) -> list[Translation]:
        """Return the presentation as one Markdown translation.

        Slides are rendered in order as ``## Slide N`` sections; empty slides are
        skipped. A presentation with no extractable text yields a single empty
        translation (consistent with the other office translators).
        """
        presentation = Presentation(io.BytesIO(_as_deck(document.content)))
        sections: list[str] = []
        # enumerate from 1 so slide numbers match what the user sees in PowerPoint.
        for index, slide in enumerate(presentation.slides, start=1):
            section = _slide_markdown(index, slide)
            if section is not None:
                sections.append(section)
        return [make_translation(document, "\n\n".join(sections))]


# Module-level instance selected by --translator pptx.
translator: Pptx = Pptx()
