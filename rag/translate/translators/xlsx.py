"""The ``xlsx`` translator (PLAN §5.4).

Converts an Excel workbook to Markdown: one section per worksheet, each rendered
as a Markdown table of its cells. Handles both the modern OOXML formats
(``.xlsx``/``.xltx`` and their macro-enabled ``.xlsm``/``.xltm`` siblings, via
``openpyxl``) and the legacy binary ``.xls`` (BIFF) format (via ``xlrd``); the two
paths feed the same table renderer. The format is detected from the content — a
ZIP local-file header means OOXML, otherwise it is treated as legacy BIFF.
"""

from __future__ import annotations

import io

import openpyxl
import xlrd

from ..base import Document, Translation, Translator, make_translation

# The first two bytes of every ZIP archive (``PK``); OOXML workbooks are ZIPs,
# legacy .xls (an OLE2 compound file) is not, so this cleanly separates the paths.
_ZIP_MAGIC: bytes = b"PK"


def _md_safe(text: str) -> str:
    """Escape/flatten a cell string so it stays inside one Markdown table cell.

    Pipes are escaped (they are the table delimiter) and any newlines are flattened
    so one spreadsheet row stays on a single Markdown line.
    """
    text = text.replace("|", "\\|")
    return text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")


def _cell(value: object) -> str:
    """Render one openpyxl cell value as Markdown-table-safe text.

    ``None`` (empty cell) becomes an empty string; every other value is stringified
    and passed through :func:`_md_safe`.
    """
    if value is None:
        return ""
    return _md_safe(str(value))


def _number(value: float) -> str:
    """Render a numeric cell without a spurious trailing ``.0``.

    Excel stores every number as a float, so an integer like ``3`` arrives as
    ``3.0``. Whole numbers are rendered without the fractional part to match the
    openpyxl path (which yields a genuine ``int``); non-whole numbers keep their
    full ``str`` representation.
    """
    return str(int(value)) if value == int(value) else str(value)


def _xls_cell(cell: xlrd.sheet.Cell, datemode: int) -> str:
    """Render one legacy-.xls (xlrd) cell as Markdown-table-safe text.

    xlrd exposes a typed cell; we map each BIFF cell type to a readable string:
    numbers drop the trailing ``.0`` for whole values, dates become ISO strings,
    booleans become ``TRUE``/``FALSE`` and error cells their error text. Empty and
    blank cells become an empty string.
    """
    ctype = cell.ctype
    if ctype == xlrd.XL_CELL_NUMBER:
        return _md_safe(_number(cell.value))
    if ctype == xlrd.XL_CELL_DATE:
        # Convert the Excel serial date to a datetime; drop a zero time-of-day so a
        # pure date does not carry a misleading ``00:00:00``.
        dt = xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        text = dt.date().isoformat() if dt.time() == dt.min.time() else dt.isoformat(" ")
        return _md_safe(text)
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if ctype == xlrd.XL_CELL_ERROR:
        # xlrd maps the raw error code to its Excel spelling (e.g. ``#DIV/0!``).
        return xlrd.error_text_from_code.get(cell.value, "#ERR")
    if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    return _md_safe(str(cell.value))


def _sheet_to_markdown(title: str, rows: list[list[str]]) -> str | None:
    """Render one worksheet's rows as a Markdown table, or ``None`` if empty.

    Trailing all-empty rows and columns are dropped; every remaining row is padded
    to a common width. The first row is used as the table header (Markdown tables
    require a header + separator, even when the sheet has no real header row).
    """
    # Drop trailing all-empty rows.
    while rows and all(cell == "" for cell in rows[-1]):
        rows.pop()
    if not rows:
        return None
    width = max(len(row) for row in rows)
    # Drop trailing all-empty columns (keep at least one).
    while width > 1 and all(len(row) < width or row[width - 1] == "" for row in rows):
        width -= 1
    # Normalise every row to the common width.
    norm = [(row + [""] * width)[:width] for row in rows]
    lines = [
        f"## {title}",
        "| " + " | ".join(norm[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines += ["| " + " | ".join(row) + " |" for row in norm[1:]]
    return "\n".join(lines)


def _xlsx_sections(content: bytes) -> list[str]:
    """Render every non-empty sheet of an OOXML workbook to a Markdown section.

    Cached cell values are read (``data_only``) — formulas are not evaluated, so a
    workbook last saved by a non-Excel writer may expose ``None`` for formula cells.
    """
    # read_only streams large sheets without loading them wholesale; data_only
    # returns the last cached value rather than the formula string.
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sections: list[str] = []
        for worksheet in workbook.worksheets:
            rows = [
                [_cell(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            section = _sheet_to_markdown(worksheet.title, rows)
            if section is not None:
                sections.append(section)
        return sections
    finally:
        # read_only mode holds the underlying zip open until explicitly closed.
        workbook.close()


def _xls_sections(content: bytes) -> list[str]:
    """Render every non-empty sheet of a legacy .xls (BIFF) workbook to Markdown.

    xlrd reads the whole book from the in-memory bytes; each sheet's cells are
    stringified via :func:`_xls_cell` (which needs the book's ``datemode`` to
    interpret serial dates) and fed to the shared table renderer.
    """
    book = xlrd.open_workbook(file_contents=content)
    sections: list[str] = []
    for sheet in book.sheets():
        rows = [
            [_xls_cell(sheet.cell(r, c), book.datemode) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        section = _sheet_to_markdown(sheet.name, rows)
        if section is not None:
            sections.append(section)
    return sections


class Xlsx(Translator):
    """Converts an Excel workbook (OOXML or legacy .xls) to Markdown tables."""

    def translate(self, document: Document) -> list[Translation]:
        """Return the workbook as one Markdown translation (a table per sheet).

        The format is sniffed from the content: a ZIP header means an OOXML
        workbook (openpyxl), anything else is treated as legacy BIFF .xls (xlrd).
        Empty sheets are skipped; sheets are separated by blank lines.
        """
        content = document.content
        # ZIP magic ⇒ OOXML (.xlsx/.xlsm/…); otherwise legacy OLE2 BIFF (.xls).
        if content[:2] == _ZIP_MAGIC:
            sections = _xlsx_sections(content)
        else:
            sections = _xls_sections(content)
        return [make_translation(document, "\n\n".join(sections))]


# Module-level instance selected by --translator xlsx.
translator: Xlsx = Xlsx()
