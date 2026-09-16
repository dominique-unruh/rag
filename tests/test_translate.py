"""Tests for module 2 (translation): translators, discovery, store, pipeline."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from rag.blobs import BlobStore
from rag.config import Config
from rag.db import Database
from rag.extract.store import SqliteStore
from rag.translate.base import Document, Translation
from rag.translate.discovery import TranslatorError, list_translators, load_translator
from rag.translate.pipeline import TranslateOptions, TranslationError, _validate, run_translation
from rag.translate.store import SourceDocument, TranslationStore
from rag.translate.translators.dispatch import translator as dispatch
from rag.translate.translators.text import translator as text_translator

# A multipart email: plain body + one text attachment, with an encoded subject.
_EMAIL = (
    "From: A <a@x>\r\n"
    "To: B <b@y>\r\n"
    "Subject: =?utf-8?q?Gr=C3=BC=C3=9Fe?=\r\n"
    "MIME-Version: 1.0\r\n"
    "Content-Type: multipart/mixed; boundary=BND\r\n\r\n"
    "--BND\r\nContent-Type: text/plain; charset=utf-8\r\n\r\nHello body.\r\n"
    "--BND\r\nContent-Type: text/plain\r\n"
    "Content-Disposition: attachment; filename=note.txt\r\n\r\nAttached text.\r\n"
    "--BND--\r\n"
).encode()


# --- translators ---------------------------------------------------------
def test_text_translator_html_and_passthrough() -> None:
    """HTML becomes Markdown; plain text and Markdown pass through."""
    html = Document("d", [], b"<h1>Hi</h1><p>yo</p>", {"mime-type": "text/html"}, {})
    assert text_translator.translate(html)[0].text == "# Hi\n\nyo"
    plain = Document("d", [], b"just text", {"mime-type": "text/plain"}, {})
    assert plain.mime_type == "text/plain"
    assert text_translator.translate(plain)[0].text == "just text"


def test_email_translator_body_and_attachment() -> None:
    """Email yields a body part (decoded headers) + attachment part via dispatch."""
    doc = Document("m1", [], _EMAIL, {"mime-type": "message/rfc822"}, {})
    parts = dispatch.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    # Body: decoded subject prepended, body text present.
    body = by_sub[("body",)]
    assert "**Subject:** Grüße" in body.text
    assert "Hello body." in body.text
    # Attachment routed through dispatch -> text translator.
    assert by_sub[("attachment:1",)].text == "Attached text."
    # Every translation keeps the source doc_id.
    assert all(p.doc_id == "m1" for p in parts)


def test_dispatch_tags_debug_info() -> None:
    """Dispatch records where it routed each translation (nested trace)."""
    doc = Document("m1", [], _EMAIL, {"mime-type": "message/rfc822"}, {})
    by_sub = {tuple(p.sub_doc_id): p for p in dispatch.translate(doc)}
    # Body routed straight to Email.
    assert by_sub[("body",)].debug_info == "Dispatched to Email"
    # Attachment: inner dispatch -> Text, then outer dispatch -> Email.
    assert by_sub[("attachment:1",)].debug_info == "Dispatched to Text\nDispatched to Email"


def test_pdf_translator_returns_single_part() -> None:
    """The PDF translator returns exactly one translation for a valid PDF."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    from rag.translate.translators.pdf import translator as pdf_translator

    doc = Document("p", [], buf.getvalue(), {"mime-type": "application/pdf"}, {})
    result = pdf_translator.translate(doc)
    assert len(result) == 1 and result[0].doc_id == "p"


def _minimal_docx(body_xml: str) -> bytes:
    """Build a minimal valid .docx (OOXML zip) whose body is ``body_xml``."""
    import zipfile

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body_xml}</w:body></w:document>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
    return buf.getvalue()


def test_docx_translator_converts_to_markdown() -> None:
    """The docx translator maps Word styles to Markdown (headings, bold)."""
    from rag.translate.translators.docx import translator as docx_translator

    body = (
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>My Title</w:t></w:r></w:p>'
        "<w:p><w:r><w:t>Hello </w:t></w:r>"
        "<w:r><w:rPr><w:b/></w:rPr><w:t>bold</w:t></w:r>"
        "<w:r><w:t> world.</w:t></w:r></w:p>"
    )
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    doc = Document("dx", [], _minimal_docx(body), {"mime-type": mime}, {})
    result = docx_translator.translate(doc)
    assert len(result) == 1 and result[0].doc_id == "dx"
    assert result[0].text == "# My Title\n\nHello **bold** world."


def test_dispatch_routes_docx_and_dotx() -> None:
    """Dispatch routes docx, dotx (template) and macro-enabled (.docm/.dotm) to Docx."""
    body = "<w:p><w:r><w:t>plain para</w:t></w:r></w:p>"
    for mime in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.template",
        "application/vnd.ms-word.document.macroenabled.12",
        "application/vnd.ms-word.template.macroenabled.12",
    ):
        doc = Document("d", [], _minimal_docx(body), {"mime-type": mime}, {})
        parts = dispatch.translate(doc)
        assert len(parts) == 1
        assert parts[0].text == "plain para"
        assert parts[0].debug_info == "Dispatched to Docx"


def _make_xlsx() -> bytes:
    """Build an in-memory .xlsx: a data sheet, an empty sheet, and a sparse one."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty"])
    ws.append(["Widget | A", 3])  # pipe must be escaped
    ws.append(["Gadget", None])   # empty cell
    wb.create_sheet("Empty")      # no cells -> skipped
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_translator_renders_markdown_tables() -> None:
    """The xlsx translator emits one Markdown table per non-empty sheet."""
    from rag.translate.translators.xlsx import translator as xlsx_translator

    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    doc = Document("xl", [], _make_xlsx(), {"mime-type": mime}, {})
    result = xlsx_translator.translate(doc)
    assert len(result) == 1 and result[0].doc_id == "xl"
    text = result[0].text
    assert text == (
        "## Sales\n"
        "| Item | Qty |\n"
        "| --- | --- |\n"
        "| Widget \\| A | 3 |\n"
        "| Gadget |  |"
    )
    # The empty sheet contributes nothing.
    assert "Empty" not in text


def test_dispatch_routes_xlsx() -> None:
    """Dispatch routes xlsx, xltx (template) and macro-enabled (.xlsm/.xltm) to Xlsx."""
    data = _make_xlsx()
    for mime in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.template",
        "application/vnd.ms-excel.sheet.macroenabled.12",
        "application/vnd.ms-excel.template.macroenabled.12",
    ):
        doc = Document("xl", [], data, {"mime-type": mime}, {})
        parts = dispatch.translate(doc)
        assert len(parts) == 1
        assert parts[0].debug_info == "Dispatched to Xlsx"
        assert parts[0].text.startswith("## Sales")


def _make_pptx() -> bytes:
    """Build an in-memory .pptx: a text slide (+notes), a table slide, an empty one."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]  # the "Blank" layout: no stray placeholders

    # Slide 1: a multi-paragraph textbox plus speaker notes.
    s1 = prs.slides.add_slide(blank)
    tf = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
    tf.text = "Quarterly Review"       # first paragraph
    tf.add_paragraph().text = "Revenue up"
    tf.add_paragraph().text = "Costs down"
    s1.notes_slide.notes_text_frame.text = "remember to smile"

    # Slide 2: a 2x2 table (first row is the header).
    s2 = prs.slides.add_slide(blank)
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Users"
    table.cell(1, 1).text = "42"

    # Slide 3: no shapes at all -> skipped in the output.
    prs.slides.add_slide(blank)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_translator_renders_slides() -> None:
    """The pptx translator emits a section per non-empty slide, with notes + tables."""
    from rag.translate.translators.pptx import translator as pptx_translator

    mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    doc = Document("pp", [], _make_pptx(), {"mime-type": mime}, {})
    result = pptx_translator.translate(doc)
    assert len(result) == 1 and result[0].doc_id == "pp"
    assert result[0].text == (
        "## Slide 1\n\n"
        "Quarterly Review\nRevenue up\nCosts down\n\n"
        "### Notes\n\n"
        "remember to smile\n\n"
        "## Slide 2\n\n"
        "| Metric | Value |\n"
        "| --- | --- |\n"
        "| Users | 42 |"
    )
    # The empty third slide contributes no section.
    assert "## Slide 3" not in result[0].text


def _retype_as_template(deck: bytes) -> bytes:
    """Re-pack a .pptx deck as a .potx template (rewrite its main content type).

    Produces bytes whose ``[Content_Types].xml`` uses the template spelling that
    python-pptx rejects, so the translator's normalisation path is exercised for
    real (not just via the routing mime).
    """
    import zipfile

    deck_ct = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    )
    template_ct = (
        "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
    )
    src = zipfile.ZipFile(io.BytesIO(deck))
    ct = src.read("[Content_Types].xml").decode("utf-8").replace(deck_ct, template_ct)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = ct.encode("utf-8") if item.filename == "[Content_Types].xml" else src.read(item.filename)
            dst.writestr(item, data)
    return out.getvalue()


def test_pptx_translator_handles_template() -> None:
    """A .potx (template content type) is normalised to a deck and translated."""
    from rag.translate.translators.pptx import translator as pptx_translator

    mime = "application/vnd.openxmlformats-officedocument.presentationml.template"
    doc = Document("pt", [], _retype_as_template(_make_pptx()), {"mime-type": mime}, {})
    result = pptx_translator.translate(doc)
    assert len(result) == 1
    assert result[0].text.startswith("## Slide 1")


def _retype_as_macro(deck: bytes) -> bytes:
    """Re-pack a .pptx deck as a macro-enabled .pptm (rewrite its main content type).

    The macro-enabled presentation uses the ``ms-powerpoint`` namespace with the
    ``macroEnabled`` spelling for its main part — another content type python-pptx
    rejects — so this exercises the ``_as_deck`` normalisation for that variant.
    """
    import zipfile

    deck_ct = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    )
    macro_ct = "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"
    src = zipfile.ZipFile(io.BytesIO(deck))
    ct = src.read("[Content_Types].xml").decode("utf-8").replace(deck_ct, macro_ct)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = ct.encode("utf-8") if item.filename == "[Content_Types].xml" else src.read(item.filename)
            dst.writestr(item, data)
    return out.getvalue()


def test_pptx_translator_handles_macro_enabled() -> None:
    """A .pptm (macro-enabled content type) is normalised to a deck and translated."""
    from rag.translate.translators.pptx import translator as pptx_translator

    mime = "application/vnd.ms-powerpoint.presentation.macroenabled.12"
    doc = Document("pm", [], _retype_as_macro(_make_pptx()), {"mime-type": mime}, {})
    result = pptx_translator.translate(doc)
    assert len(result) == 1
    assert result[0].text.startswith("## Slide 1")


def test_dispatch_routes_pptx_family() -> None:
    """Dispatch routes deck/template/slideshow + macro-enabled presentation mimes to Pptx."""
    data = _make_pptx()
    for mime in (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.openxmlformats-officedocument.presentationml.template",
        "application/vnd.openxmlformats-officedocument.presentationml.slideshow",
        "application/vnd.ms-powerpoint.presentation.macroenabled.12",
        "application/vnd.ms-powerpoint.template.macroenabled.12",
        "application/vnd.ms-powerpoint.slideshow.macroenabled.12",
    ):
        doc = Document("pp", [], data, {"mime-type": mime}, {})
        parts = dispatch.translate(doc)
        assert len(parts) == 1
        assert parts[0].debug_info == "Dispatched to Pptx"
        assert parts[0].text.startswith("## Slide 1")


def _make_zip() -> bytes:
    """Build an in-memory .zip: a text file, a nested dir, and a directory entry."""
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("readme.txt", "hello from zip")
        archive.writestr("docs/note.md", "# Heading\n\nbody")
        archive.writestr("emptydir/", "")  # a directory entry: must be skipped
    return buf.getvalue()


def test_archive_translator_unpacks_zip() -> None:
    """The archive translator dispatches each contained file (skipping dirs)."""
    from rag.translate.translators.archive import translator as archive_translator

    doc = Document("z", [], _make_zip(), {"mime-type": "application/zip"}, {})
    parts = archive_translator.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    # A leading TOC part plus the two files; the directory entry produced nothing.
    assert set(by_sub) == {("toc",), ("readme.txt",), ("docs/note.md",)}
    assert by_sub[("readme.txt",)].text == "hello from zip"
    assert by_sub[("docs/note.md",)].text == "# Heading\n\nbody"
    # The TOC is emitted first and lists every file.
    assert parts[0].sub_doc_id == ["toc"]
    assert "readme.txt" in parts[0].text and "docs/note.md" in parts[0].text


def test_archive_toc_survives_untranslatable_members() -> None:
    """Even when every member is untranslatable, the TOC still yields a part."""
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        # A binary (NUL-containing) blob dispatch cannot translate.
        archive.writestr("image.bin", b"\x00\x01\x02\x03")
    doc = Document("b", [], buf.getvalue(), {"mime-type": "application/zip"}, {})
    parts = dispatch.translate(doc)
    # Only the TOC remains, and it names the otherwise-lost binary member.
    assert [p.sub_doc_id for p in parts] == [["toc"]]
    assert "image.bin" in parts[0].text


def test_archive_tolerates_a_corrupt_member() -> None:
    """A member that crashes its translator is skipped, not fatal to the archive."""
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("broken.pdf", b"%PDF-1.4 truncated garbage")  # crashes pypdf
        archive.writestr("ok.txt", "still here")
    doc = Document("c", [], buf.getvalue(), {"mime-type": "application/zip"}, {})
    parts = dispatch.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    # The good file and the TOC survive; the corrupt PDF is dropped.
    assert set(by_sub) == {("toc",), ("ok.txt",)}
    assert by_sub[("ok.txt",)].text == "still here"
    assert "broken.pdf" in by_sub[("toc",)].text  # still listed in the TOC


def test_dispatch_routes_zip() -> None:
    """Dispatch routes a zip mime type to Archive, extending sub_doc_ids and trace."""
    doc = Document("z", [], _make_zip(), {"mime-type": "application/zip"}, {})
    parts = dispatch.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    assert set(by_sub) == {("toc",), ("readme.txt",), ("docs/note.md",)}
    # Inner dispatch -> Text, then outer dispatch -> Archive.
    assert by_sub[("readme.txt",)].debug_info == "Dispatched to Text\nDispatched to Archive"


def test_archive_translator_unpacks_tar_gz() -> None:
    """A gzip-compressed tar is unpacked into its file members (not one gzip blob)."""
    import io as _io
    import tarfile

    buf = _io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w:gz") as archive:
        payload = b"tarred text"
        info = tarfile.TarInfo("a/b.txt")
        info.size = len(payload)
        archive.addfile(info, _io.BytesIO(payload))
    doc = Document("t", [], buf.getvalue(), {"mime-type": "application/gzip"}, {})
    parts = dispatch.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    assert set(by_sub) == {("toc",), ("a/b.txt",)}
    assert by_sub[("a/b.txt",)].text == "tarred text"
    assert "a/b.txt" in by_sub[("toc",)].text


def test_archive_translator_unpacks_gzip_stream() -> None:
    """A standalone gzip stream yields one part named by the stripped filename."""
    import gzip

    doc = Document(
        "g", [], gzip.compress(b"plain gz body"),
        {"mime-type": "application/gzip", "filename": "notes.txt.gz"}, {},
    )
    parts = dispatch.translate(doc)
    assert len(parts) == 1
    assert parts[0].sub_doc_id == ["notes.txt"]  # ".gz" stripped
    assert parts[0].text == "plain gz body"


def test_archive_translator_nested_zip() -> None:
    """A zip inside a zip is unpacked recursively via dispatch."""
    import zipfile

    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as archive:
        archive.writestr("deep.txt", "deep content")
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as archive:
        archive.writestr("inner.zip", inner.getvalue())
    doc = Document("n", [], outer.getvalue(), {"mime-type": "application/zip"}, {})
    parts = dispatch.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    # Outer TOC, the inner archive's own TOC, and the deep file.
    assert set(by_sub) == {("toc",), ("inner.zip", "toc"), ("inner.zip", "deep.txt")}
    assert by_sub[("inner.zip", "deep.txt")].text == "deep content"


def test_archive_translator_unpacks_7z() -> None:
    """A 7-Zip archive is unpacked and each file member dispatched."""
    import py7zr

    buf = io.BytesIO()
    with py7zr.SevenZipFile(buf, "w") as archive:
        archive.writestr("hello 7z", "a.txt")
        archive.writestr("# H\n\nbody", "docs/b.md")
    doc = Document("s", [], buf.getvalue(), {"mime-type": "application/x-7z-compressed"}, {})
    parts = dispatch.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    assert set(by_sub) == {("toc",), ("a.txt",), ("docs/b.md",)}
    assert by_sub[("a.txt",)].text == "hello 7z"
    assert by_sub[("a.txt",)].debug_info == "Dispatched to Text\nDispatched to Archive"


# A tiny real RAR5 archive (base64): files ``a.txt`` = "hello rar" and
# ``docs/b.md`` = "# H\n\nbody". Committed as bytes because creating a .rar needs
# the proprietary ``rar`` binary; reading it needs only unrar/bsdtar/unar.
_RAR_B64 = (
    "UmFyIRoHAQAzkrXlCgEFBgAFAQGAgADL7qgeIwIDC4kABIkApIMCF8lRhIAAAQVhLnR4dAoD"
    "E3WJqmrsOj8LaGVsbG8gcmFyOaoYXScCAwuJAASJAKSDAukcIMKAAAEJZG9jcy9iLm1kCgMT"
    "dYmqauw6PwsjIEgKCmJvZHkdd1ZRAwUEAA=="
)


def test_archive_translator_unpacks_rar() -> None:
    """A RAR archive is unpacked and each file member dispatched."""
    import base64

    doc = Document("r", [], base64.b64decode(_RAR_B64), {"mime-type": "application/vnd.rar"}, {})
    parts = dispatch.translate(doc)
    by_sub = {tuple(p.sub_doc_id): p for p in parts}
    assert set(by_sub) == {("toc",), ("a.txt",), ("docs/b.md",)}
    assert by_sub[("a.txt",)].text == "hello rar"
    assert by_sub[("docs/b.md",)].text == "# H\n\nbody"


def test_dispatch_text_subtype_identity() -> None:
    """An unsupported text/* type (e.g. .tex) is translated as-is via Text."""
    tex = "\\section{Hi}\nbody\n"
    doc = Document("t", [], tex.encode(), {"mime-type": "text/x-tex"}, {})
    parts = dispatch.translate(doc)
    assert len(parts) == 1
    assert parts[0].text == tex  # identity passthrough, no conversion
    assert parts[0].debug_info == "Dispatched to Text"


def test_dispatch_unknown_but_texty_identity() -> None:
    """An unrecognised mime whose bytes look like text gets an identity translation."""
    doc = Document("u", [], b"just some plain words\nline two\n",
                   {"mime-type": "application/x-unknown"}, {})
    parts = dispatch.translate(doc)
    assert len(parts) == 1
    assert "plain words" in parts[0].text
    assert parts[0].debug_info == "Dispatched to Text"


def test_dispatch_unknown_mime_yields_nothing() -> None:
    """An unhandled, binary (NUL-containing) mime type produces no translation."""
    doc = Document("d", [], b"\x00\x01", {"mime-type": "application/octet-stream"}, {})
    assert dispatch.translate(doc) == []


def test_dispatch_ignored_mime_yields_nothing() -> None:
    """A deliberately-ignored mime (e.g. S/MIME signature) is routed to `ignore`."""
    doc = Document("d", [], b"\x00\x01", {"mime-type": "application/x-pkcs7-signature"}, {})
    assert dispatch.translate(doc) == []
    assert load_translator("ignore").translate(doc) == []


# --- discovery -----------------------------------------------------------
def test_discovery() -> None:
    """Built-in translators are discoverable; unknown names raise."""
    assert {"dispatch", "email", "text", "pdf", "archive", "pptx"} <= set(list_translators())
    assert load_translator("text") is text_translator
    with pytest.raises(TranslatorError):
        load_translator("does-not-exist")


# --- pipeline invariants -------------------------------------------------
def _source() -> SourceDocument:
    """A minimal source document for validation tests."""
    return SourceDocument("email", "m1", 1, b"", {}, {})


def test_validate_rejects_wrong_doc_id() -> None:
    """A translation with a mismatched doc_id is rejected."""
    with pytest.raises(TranslationError):
        _validate(_source(), [Translation("OTHER", ["body"], "t", {})])


def test_validate_rejects_duplicate_sub_doc_id() -> None:
    """Two translations sharing a sub_doc_id are rejected."""
    dup = [Translation("m1", ["body"], "a", {}), Translation("m1", ["body"], "b", {})]
    with pytest.raises(TranslationError):
        _validate(_source(), dup)


# --- store ---------------------------------------------------------------
def test_store_select_and_write(database: Database, blob_store: BlobStore) -> None:
    """select_documents finds untranslated/stale docs; write_set stores an active set."""
    src = SqliteStore(database.conn, blob_store, "email")
    src.upsert("m1", _EMAIL, {"mime-type": "message/rfc822", "subject": "Hi"})
    store = TranslationStore(database.conn, blob_store)

    # Untranslated -> selected.
    assert [d.doc_id for d in store.select_documents()] == ["m1"]

    set_id = store.write_set(
        source_type="email", doc_id="m1", source_version=1,
        translator_name="dispatch", git_rev="rev",
        parts=[(["body"], "A", {}), (["attachment:1"], "B", {})],
    )
    conn = database.conn
    assert conn.execute("SELECT COUNT(*) FROM translation_parts WHERE set_id=?", (set_id,)).fetchone()[0] == 2
    keys = sorted(r[0] for r in conn.execute("SELECT part_key FROM translation_parts WHERE set_id=?", (set_id,)))
    assert keys == ['["attachment:1"]', '["body"]']

    # Up to date now -> not selected by default, but --force still yields it.
    assert list(store.select_documents()) == []
    assert [d.doc_id for d in store.select_documents(force=True)] == ["m1"]

    # A newer set deactivates the previous one (at most one active).
    set_id2 = store.write_set(
        source_type="email", doc_id="m1", source_version=1,
        translator_name="dispatch", git_rev="rev", parts=[(["body"], "newer", {})],
    )
    active = conn.execute("SELECT id FROM translation_sets WHERE is_active=1").fetchall()
    assert len(active) == 1 and active[0][0] == set_id2


def test_store_selects_stale(database: Database, blob_store: BlobStore) -> None:
    """A document whose active set is from an older version is re-selected."""
    src = SqliteStore(database.conn, blob_store, "email")
    src.upsert("m1", _EMAIL, {"mime-type": "message/rfc822"})
    store = TranslationStore(database.conn, blob_store)
    store.write_set(source_type="email", doc_id="m1", source_version=1,
                    translator_name="dispatch", git_rev="rev", parts=[(["body"], "A", {})])
    assert list(store.select_documents()) == []
    # Change the source -> version bumps to 2 -> active set (v1) is now stale.
    src.upsert("m1", _EMAIL + b"more", {"mime-type": "message/rfc822"})
    assert [d.doc_id for d in store.select_documents()] == ["m1"]


# --- pipeline end-to-end (inspection mode; no git/DB write needed) -------
def test_run_translation_inspect(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """run_translation --inspect translates a seeded email and prints it."""
    storage = tmp_path / "storage"
    db = Database.open(storage / "metadata.db")
    blobs = BlobStore(storage / "blobs")
    SqliteStore(db.conn, blobs, "email").upsert("m1", _EMAIL, {"mime-type": "message/rfc822"})
    config = Config({"storage": {"dir": str(storage)}}, tmp_path / "config.yaml")

    processed = run_translation(config, TranslateOptions(inspect=True), db=db)

    assert processed == 1
    out = capsys.readouterr().out
    # The body (with decoded subject) and the attachment both appear.
    assert "**Subject:** Grüße" in out
    assert "Attached text." in out
    # Nothing was written to the DB in inspection mode.
    assert db.conn.execute("SELECT COUNT(*) FROM translation_sets").fetchone()[0] == 0
    db.close()
